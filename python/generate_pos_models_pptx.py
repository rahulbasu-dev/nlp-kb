"""Generate educational PPTX for Advanced POS Tagging Models.

Covers: Forward Algorithm, Viterbi, MEMM, CRF, BiLSTM-CRF, BERT.
Reuses slide builder helpers from generate_topic_pptx.py.
"""

import sys
import os
sys.path.insert(0, os.path.dirname(__file__))

from generate_topic_pptx import (
    _new_prs, add_title_slide, add_section_slide, add_content_slide,
    add_two_column_slide, add_formula_slide, add_table_slide,
    add_key_concept_slide,
    ACCENT_BLUE, ACCENT_GREEN, ACCENT_ORANGE, ACCENT_RED,
)
from pptx.dml.color import RGBColor

ACCENT_SKY = RGBColor(0x0E, 0xA5, 0xE9)


def build_pos_models(prs):
    add_title_slide(prs, "Advanced POS Tagging Models",
                    "From the Forward Algorithm & Viterbi\n"
                    "Through MEMMs & CRFs to Neural Taggers",
                    ACCENT_SKY)

    # ── Section 1: Intuition ──
    add_section_slide(prs, "1. Beyond HMMs: Why We Need More", ACCENT_SKY)

    add_content_slide(prs, "POS Tagging Model Evolution", [
        "Rule-based: Hand-crafted rules (brittle but interpretable)",
        "Statistical (HMM): Generative model, Viterbi decoding, ~95% accuracy",
        "Discriminative (MEMM): Rich features, P(tags|words) directly, ~96%",
        "Discriminative (CRF): Global normalization, solves label bias, ~97%",
        "Neural (BiLSTM-CRF): Learned features + sequence optimization, ~97.5%",
        "Pre-trained (BERT): Deep bidirectional context, ~98%",
    ], ACCENT_SKY)

    add_content_slide(prs, "HMM Limitations", [
        "Independence assumption:",
        "  - Observed word depends ONLY on current tag",
        "  - Cannot use neighboring words, suffixes, word shape, etc.",
        "",
        "Generative vs. Discriminative:",
        "  - HMM models joint P(words, tags), but we only need P(tags|words)",
        "  - Discriminative models are more direct and often more accurate",
        "",
        "Limited context:",
        "  - Standard HMMs: only bigram tag context (previous tag)",
        "  - Cannot look at word two positions back or ahead",
        "",
        'Analogy: Reading "bank" left-to-right vs. also seeing "...of the river"',
    ], ACCENT_SKY)

    # ── Section 2: Forward Algorithm ──
    add_section_slide(prs, "2. The Forward Algorithm", ACCENT_SKY)

    add_key_concept_slide(prs, "Problem: Computing P(O|lambda)",
        "Given an HMM and an observation sequence, what is the total probability "
        "of seeing that sequence? The naive approach enumerates all N^T state "
        "sequences -- completely intractable!",
        [
            "Forward variable alpha_t(j) = P(o_1...o_t, q_t=j | lambda)",
            "Captures probability of seeing first t observations AND being in state j at time t",
            "Dynamic programming: reuse intermediate results",
            "Complexity: O(N^2 * T) instead of O(N^T)",
        ], ACCENT_SKY)

    add_formula_slide(prs, "Forward Algorithm: Three Steps",
        [
            "Init:  alpha_1(j) = pi_j * b_j(o_1)",
            "Recursion:  alpha_t(j) = [SUM_i alpha_{t-1}(i) * a_ij] * b_j(o_t)",
            "Termination:  P(O|lambda) = SUM_j alpha_T(j)",
        ],
        [
            "Start prob x emission prob for each state",
            "Sum over all previous states (key difference from Viterbi: SUM not MAX)",
            "Total probability = sum of all forward variables at final time step",
        ], ACCENT_SKY)

    add_content_slide(prs, "Worked Example: Weather/Ice Cream HMM", [
        "States: Hot (H), Cold (C);  pi = [0.8, 0.2]",
        "Transitions: a_HH=0.6, a_HC=0.4, a_CH=0.5, a_CC=0.5",
        "Emissions: b_H(3)=0.4, b_C(3)=0.1, b_H(1)=0.2, b_C(1)=0.5",
        "",
        "Observation: O = (3, 1, 3)",
        "",
        "t=1: alpha_1(H)=0.8*0.4=0.32,  alpha_1(C)=0.2*0.1=0.02",
        "t=2: alpha_2(H)=[0.32*0.6+0.02*0.5]*0.2=0.0404",
        "  -  alpha_2(C)=[0.32*0.4+0.02*0.5]*0.5=0.069",
        "t=3: alpha_3(H)=[0.0404*0.6+0.069*0.5]*0.4=0.02350",
        "  -  alpha_3(C)=[0.0404*0.4+0.069*0.5]*0.1=0.00507",
        "",
        "P(O|lambda) = 0.02350 + 0.00507 = 0.02856",
    ], ACCENT_SKY)

    # ── Section 3: Viterbi ──
    add_section_slide(prs, "3. Viterbi Decoding (Deep Dive)", ACCENT_SKY)

    add_two_column_slide(prs, "Forward vs. Viterbi: Sum vs. Max",
        "Forward Algorithm",
        [
            "Goal: P(O|lambda) -- total likelihood",
            "Operation: SUM over all paths",
            "No backtracking needed",
            "Use: model evaluation, training (EM)",
            "Complexity: O(N^2 * T)",
        ],
        "Viterbi Algorithm",
        [
            "Goal: Best state sequence Q*",
            "Operation: MAX over paths",
            "Backpointers required for traceback",
            "Use: decoding (actual tagging)",
            "Complexity: O(N^2 * T)",
        ])

    add_formula_slide(prs, "Viterbi Formulas",
        [
            "v_t(j) = max_i [v_{t-1}(i) * a_ij] * b_j(o_t)",
            "bp_t(j) = argmax_i [v_{t-1}(i) * a_ij]",
            "Log space: log v_t(j) = max_i [log v_{t-1}(i) + log a_ij] + log b_j(o_t)",
        ],
        [
            "Replace SUM with MAX -- find the single best incoming path",
            "Record WHICH state gave the max (backpointer for traceback)",
            "Log space: multiplications become additions (prevents underflow)",
        ], ACCENT_SKY)

    add_content_slide(prs, "Worked Example: Viterbi on O=(3,1,3)", [
        "Same HMM as Forward Algorithm example",
        "",
        "t=1: v_1(H)=0.32,  v_1(C)=0.02",
        "t=2: v_2(H)=max(0.32*0.6, 0.02*0.5)*0.2 = 0.192*0.2 = 0.0384, bp=H",
        "  -  v_2(C)=max(0.32*0.4, 0.02*0.5)*0.5 = 0.128*0.5 = 0.064, bp=H",
        "t=3: v_3(H)=max(0.0384*0.6, 0.064*0.5)*0.4 = 0.032*0.4 = 0.0128, bp=C",
        "  -  v_3(C)=max(0.0384*0.4, 0.064*0.5)*0.1 = 0.032*0.1 = 0.0032, bp=C",
        "",
        "Best final state: H (0.0128 > 0.0032)",
        "Backtrace: H <- C <- H",
        "Best path: H -> C -> H",
        "",
        "Warning: Greedy (most likely at each step) may differ from Viterbi (globally optimal)!",
    ], ACCENT_SKY)

    # ── Section 4: MEMM ──
    add_section_slide(prs, "4. Maximum Entropy Markov Models (MEMMs)", ACCENT_SKY)

    add_two_column_slide(prs, "Generative vs. Discriminative",
        "Generative (HMM)",
        [
            "Models: P(words, tags)",
            "Features: Only word identity",
            "Independence: Words independent given tags",
            "Training: Count-based (MLE)",
            "Unknown words: Smoothing hacks required",
        ],
        "Discriminative (MEMM)",
        [
            "Models: P(tags | words) directly",
            "Features: Arbitrary (suffixes, shape, neighbors)",
            "No independence assumption needed",
            "Training: Logistic regression / MaxEnt",
            "Unknown words: Handled by features naturally",
        ])

    add_formula_slide(prs, "MEMM: Local Classifier at Each Position",
        [
            "P(t_i | t_{i-1}, o) = softmax(w . f(t_i, t_{i-1}, o))",
            "= exp(w . f(t_i, t_{i-1}, o)) / SUM_{t'} exp(w . f(t', t_{i-1}, o))",
        ],
        [
            "Softmax (maximum entropy) classifier at each position",
            "f is a feature vector, w is learned weights. Normalization is LOCAL (per position)",
        ], ACCENT_SKY)

    add_content_slide(prs, "MEMM Feature Engineering", [
        "The power of MEMMs: rich, arbitrary features",
        "",
        "Common POS tagging features:",
        "  - Current word: w_i = 'running'",
        "  - Previous tag: t_{i-1} = VBD",
        "  - Suffixes: -ing, -ed, -ly, -tion, -ness, -able",
        "  - Prefixes: un-, re-, pre-",
        "  - Word shape: Xx (capitalized), XXXX (all caps)",
        "  - Capitalization: is_capitalized, is_all_caps",
        "  - Neighboring words: w_{i-1}, w_{i+1}",
        "  - Contains digit, hyphen, etc.",
        "",
        'Example: "running" with prev_tag=VBD',
        "  suffix=-ing & tag=VBG -> 1;  prev_tag=VBD & tag=VBG -> 1",
    ], ACCENT_SKY)

    add_content_slide(prs, "The Label Bias Problem", [
        "Critical flaw of MEMMs: LABEL BIAS",
        "",
        "Because normalization is LOCAL (per state), states with few",
        "outgoing transitions effectively IGNORE observations.",
        "",
        "Example: A state that can only transition to one other state",
        "  - No matter what word is observed, 100% of probability",
        "    mass goes to that single successor",
        "  - The observation is completely ignored!",
        "",
        "Root cause: local normalization forces probabilities to sum",
        "to 1 at each state, regardless of what is observed.",
        "",
        "Solution: CRFs with GLOBAL normalization",
    ], ACCENT_SKY)

    # ── Section 5: Bidirectionality & CRFs ──
    add_section_slide(prs, "5. Bidirectionality & CRFs", ACCENT_SKY)

    add_content_slide(prs, "Why Left-to-Right Isn't Enough", [
        "Right context resolves many ambiguities:",
        "",
        '"The lead paint was dangerous" -- "lead" is JJ, but only "paint" tells us',
        '"Please record the results" vs. "This is a record" -- next word disambiguates',
        '"The old man the boats" -- only "the boats" reveals "man" is a verb',
        "",
        "Bidirectional approach: run forward AND backward, combine at each position",
    ], ACCENT_SKY)

    add_formula_slide(prs, "Conditional Random Fields (CRFs)",
        [
            "P(y|x) = (1/Z(x)) * exp(SUM_t SUM_k lambda_k * f_k(y_t, y_{t-1}, x, t))",
            "Z(x) = SUM_{y'} exp(SUM_t SUM_k lambda_k * f_k(y'_t, y'_{t-1}, x, t))",
        ],
        [
            "GLOBAL normalization over the entire sequence (not per-state)",
            "Z(x) = partition function sums over ALL possible tag sequences",
        ], ACCENT_SKY)

    add_table_slide(prs, "HMM vs. MEMM vs. CRF",
        ["Property", "HMM", "MEMM", "CRF"],
        [
            ["Type", "Generative", "Discriminative", "Discriminative"],
            ["Models", "P(x, y)", "P(y_t|y_{t-1}, x)", "P(y|x)"],
            ["Normalization", "Global (generative)", "Local (per state)", "Global (over seq)"],
            ["Features", "Limited", "Rich, arbitrary", "Rich, arbitrary"],
            ["Label bias", "No", "Yes (problem!)", "No"],
            ["Accuracy", "~95%", "~96%", "~97%"],
        ], ACCENT_SKY)

    # ── Section 6: Neural Models ──
    add_section_slide(prs, "6. Neural Network Taggers", ACCENT_SKY)

    add_content_slide(prs, "From Handcrafted Features to Learned Representations", [
        "Word embeddings as input: dense vectors capturing semantic/syntactic properties",
        "",
        "Window-based feedforward tagger:",
        "  - Concatenate embeddings of target + neighbors (e.g., 5-word window)",
        "  - Feed into feedforward network, predict tag with softmax",
        "  - Simple but limited by fixed window size",
    ], ACCENT_SKY)

    add_formula_slide(prs, "BiLSTM Tagger Architecture",
        [
            "h_t^f = LSTM_forward(x_t, h_{t-1}^f)",
            "h_t^b = LSTM_backward(x_t, h_{t+1}^b)",
            "y_t = softmax(W * [h_t^f ; h_t^b] + b)",
        ],
        [
            "Forward LSTM reads left-to-right, encoding left context",
            "Backward LSTM reads right-to-left, encoding right context",
            "Concatenated hidden states capture FULL bidirectional context",
        ], ACCENT_SKY)

    add_content_slide(prs, "BiLSTM-CRF: Best of Both Worlds", [
        "BiLSTM-CRF combines:",
        "  - BiLSTM: Learns rich contextual representations automatically",
        "  - CRF layer: Enforces valid tag transitions, optimizes whole sequence",
        "",
        "The CRF uses BiLSTM outputs as emission scores",
        "and learns transition scores between tags.",
        "Decoding uses Viterbi over these combined scores.",
        "",
        "Transformer-based taggers (BERT):",
        "  - Self-attention: each word attends to ALL other words simultaneously",
        "  - Pre-trained on massive corpora, fine-tune for POS tagging",
        "  - Simply add linear classification layer on top of BERT",
    ], ACCENT_SKY)

    add_table_slide(prs, "Performance Comparison (Penn Treebank)",
        ["Model", "Accuracy", "Year", "Key Innovation"],
        [
            ["HMM", "~95.0%", "1960s-", "Probabilistic sequence model"],
            ["MEMM", "~96.0%", "2000", "Discriminative + rich features"],
            ["CRF", "~97.0%", "2001", "Global normalization"],
            ["BiLSTM-CRF", "~97.5%", "2015", "Learned features + seq opt."],
            ["BERT", "~97.9%", "2019", "Pre-trained deep bidirectional"],
        ], ACCENT_SKY)

    # ── Section 7: References ──
    add_section_slide(prs, "7. References", ACCENT_SKY)

    add_content_slide(prs, "Key References", [
        "[1] Jurafsky & Martin (2024) - Speech and Language Processing, Ch. 8",
        "  - HMM, Forward Algorithm, Viterbi, CRFs",
        "",
        "[2] McCallum, Freitag & Pereira (2000) - Maximum Entropy Markov Models",
        "  - for Information Extraction and Segmentation (ICML)",
        "",
        "[3] Lafferty, McCallum & Pereira (2001) - Conditional Random Fields",
        "  - Probabilistic Models for Segmenting and Labeling (ICML)",
        "",
        "[4] Huang, Xu & Yu (2015) - Bidirectional LSTM-CRF Models",
        "  - for Sequence Tagging (arXiv:1508.01991)",
        "",
        "[5] Devlin et al. (2019) - BERT: Pre-training of Deep Bidirectional",
        "  - Transformers for Language Understanding (NAACL-HLT)",
        "",
        "[6] Akbik, Blythe & Vollgraf (2018) - Contextual String Embeddings",
        "  - for Sequence Labeling (COLING)",
    ], ACCENT_SKY)


def main() -> None:
    """Generate POS Tagging Models PPTX."""
    prs = _new_prs()
    build_pos_models(prs)
    path = "../resources/pos_tagging_models_presentation.pptx"
    prs.save(path)
    print(f"Saved: {path}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
